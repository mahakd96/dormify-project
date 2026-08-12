"""
HTTP-level regression tests for the Excel report/export endpoints.

Background: manual smoke testing found that every Excel report endpoint
crashed with `NameError: name 'HttpResponse' is not defined` inside
`_excel_response()` (backend/api/views.py) - `HttpResponse` was used but
never imported anywhere in the file. The existing report tests
(test_reports_performance.py, and report_exports.py's own logic) all called
the report GENERATOR functions (generate_capacity_report(),
generate_manual_review_report(), ...) directly, never the actual HTTP view
functions - so `_excel_response()`, the one place the bug lived, was never
exercised by any automated test. The bug reached manual testing instead of
being caught by CI.

This module closes that gap: it calls each report endpoint through DRF's
APIClient (a real HTTP request through Django's URL routing, view,
permission classes, and response machinery), never `_excel_response()`,
the workbook-generation helpers, or the report generator functions
directly. See project-quality/performance/GROUP1_REPORT_HTTP_REGRESSION.md
for the full writeup, including proof that these tests would have failed
before the HttpResponse import fix.

Endpoints covered (all 5 users of _excel_response(), determined by
grepping backend/api/views.py and cross-referencing backend/api/urls.py -
not guessed):

    GET /api/reports/dormify-report/             (dormify_report,             IsCentralAdmin)
    GET /api/reports/student-allocation-report/  (student_allocation_report,  IsAuthenticated)
    GET /api/reports/student-actions-report/     (student_actions_report,     IsAuthenticated)
    GET /api/reports/capacity-report/             (capacity_report,           IsAuthenticated)
    GET /api/reports/manual-review-report/        (manual_review_report,      IsAuthenticated)

Run against Django's disposable TEST database (never touches the real
database):

    cd backend
    ENV_FILE=.env.test python manage.py test api.performance_tests.test_report_http_regression -v 2
"""

import re
from io import BytesIO

import openpyxl
from django.test import TestCase
from rest_framework.test import APIClient

from api.models import (
    User, Region, DormType, Building, Apartment, Room, Bed,
    Student, BedAssignment,
)

_XLSX_CONTENT_TYPE = (
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
)
_FILENAME_RE = re.compile(r'filename="([^"]+\.xlsx)"')


def _make_user(email, role, region=None):
    user = User(email=email, username=email, role=role, region=region, first_name='T', last_name='User')
    user.set_password('testpass123')
    user.save()
    return user


class ReportHttpRegressionTests(TestCase):
    """
    One minimal shared fixture (a region with a building/apartment/room/bed
    and one assigned student, plus one unassigned student) is enough for
    every report generator to run its normal code path without error -
    the actual point of these tests is the HTTP response wrapper, not the
    report content (that is already covered by
    test_reports_performance.py).
    """

    def setUp(self):
        self.region = Region.objects.create(id='g1-http-region', name='Group1 HTTP Report Region')
        self.dorm_type = DormType.objects.create(code=993, name='G1HttpDormType', region=self.region)

        self.admin = _make_user('g1-http-admin@test.com', User.Role.CENTRAL_ADMIN)
        self.boss = _make_user('g1-http-boss@test.com', User.Role.REGION_BOSS, self.region)

        building = Building.objects.create(number=1, dorm_type=self.dorm_type)
        apartment = Apartment.objects.create(
            building=building, number='1', category=Apartment.Category.MIXED,
            apartment_type=Apartment.ApartmentType.SINGLE, room_count=1,
        )
        room = Room.objects.create(apartment=apartment, name='1', capacity=2)
        bed = Bed.objects.create(room=room, label='B1')
        Bed.objects.create(room=room, label='B2')  # left unoccupied

        assigned = Student.objects.create(
            student_id='HTTP-ASSIGNED', first_name='F', last_name='L', gender='male',
            housing_type=Student.HousingType.SINGLE_MALE, accepted_dorm_type=self.dorm_type,
        )
        BedAssignment.objects.create(
            student=assigned, bed=bed, status=BedAssignment.Status.ACTIVE, assigned_by=self.admin,
        )
        Student.objects.create(
            student_id='HTTP-UNASSIGNED', first_name='F2', last_name='L2', gender='female',
            housing_type=Student.HousingType.SINGLE_FEMALE, accepted_dorm_type=self.dorm_type,
        )

    def _client_for(self, user):
        client = APIClient()
        client.force_authenticate(user)
        return client

    def _assert_valid_excel_response(self, resp):
        """
        The 5 required checks (§5 of the task): status 200, non-empty
        body, correct .xlsx Content-Type, an attachment
        Content-Disposition with an .xlsx filename, and bytes that
        openpyxl can actually parse as a workbook. Returns the parsed
        workbook so callers can do extra content spot-checks if useful.
        """
        self.assertEqual(resp.status_code, 200, resp.content[:1000])

        content = resp.content
        self.assertGreater(len(content), 0, 'response body must not be empty')

        content_type = resp.get('Content-Type', '')
        self.assertEqual(
            content_type, _XLSX_CONTENT_TYPE,
            f'unexpected Content-Type: {content_type!r}',
        )

        disposition = resp.get('Content-Disposition', '')
        self.assertIn('attachment', disposition, f'not a download: {disposition!r}')
        match = _FILENAME_RE.search(disposition)
        self.assertIsNotNone(match, f'no .xlsx filename in Content-Disposition: {disposition!r}')

        # The real proof this is a working Excel download and not just an
        # HttpResponse wrapping arbitrary/garbage bytes: openpyxl must be
        # able to load it as a genuine .xlsx workbook with at least one
        # sheet.
        wb = openpyxl.load_workbook(BytesIO(content))
        self.assertGreater(len(wb.sheetnames), 0, 'workbook has no sheets')
        return wb

    # --- 1. dormify_report — central-admin only (all-regions report) -----

    def test_dormify_report_http_returns_valid_xlsx(self):
        resp = self._client_for(self.admin).get('/api/reports/dormify-report/')
        wb = self._assert_valid_excel_response(resp)
        # Sanity: this report always builds the same fixed sheet set.
        self.assertIn('סיכום', wb.sheetnames)

    def test_dormify_report_http_rejects_non_central_admin(self):
        """
        Permission enforcement is part of the real HTTP path this module
        exists to protect - a regional user must NOT be able to pull the
        all-regions report.
        """
        resp = self._client_for(self.boss).get('/api/reports/dormify-report/')
        self.assertEqual(resp.status_code, 403)

    # --- 2. student_allocation_report — legacy alias, any authenticated --

    def test_student_allocation_report_http_returns_valid_xlsx(self):
        resp = self._client_for(self.boss).get('/api/reports/student-allocation-report/')
        self._assert_valid_excel_response(resp)

    # --- 3. student_actions_report ---------------------------------------

    def test_student_actions_report_http_returns_valid_xlsx(self):
        resp = self._client_for(self.boss).get(
            '/api/reports/student-actions-report/', {'region_id': self.region.id},
        )
        wb = self._assert_valid_excel_response(resp)
        self.assertIn('תמצית סטודנטים', wb.sheetnames)

    # --- 4. capacity_report -----------------------------------------------

    def test_capacity_report_http_returns_valid_xlsx(self):
        resp = self._client_for(self.admin).get(
            '/api/reports/capacity-report/', {'region_id': self.region.id},
        )
        wb = self._assert_valid_excel_response(resp)
        self.assertIn('תפוסה לפי חדרים', wb.sheetnames)

    # --- 5. manual_review_report -------------------------------------------

    def test_manual_review_report_http_returns_valid_xlsx(self):
        resp = self._client_for(self.boss).get('/api/reports/manual-review-report/')
        wb = self._assert_valid_excel_response(resp)
        self.assertIn('תמצית חריגים', wb.sheetnames)

    # --- all 5 endpoints, one parametrized sweep (belt-and-suspenders) ---

    def test_all_report_endpoints_return_valid_xlsx(self):
        """
        Every endpoint that ultimately calls _excel_response(), hit in one
        test for a single clear pass/fail signal covering the full set.
        """
        endpoints = [
            ('/api/reports/dormify-report/', self.admin, {}),
            ('/api/reports/student-allocation-report/', self.boss, {}),
            ('/api/reports/student-actions-report/', self.boss, {'region_id': self.region.id}),
            ('/api/reports/capacity-report/', self.admin, {'region_id': self.region.id}),
            ('/api/reports/manual-review-report/', self.boss, {}),
        ]
        for url, user, params in endpoints:
            with self.subTest(url=url, user=user.email):
                resp = self._client_for(user).get(url, params)
                self._assert_valid_excel_response(resp)
