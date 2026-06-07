"""
DORMIFY - Import Buildings from Excel
Faithful import with SSL reconnect support for Neon.

capacity logic:
  - If title contains "שניים בחדר" → capacity = 2 for every room
  - Otherwise → capacity = number of times same room_id appears in rows
"""
import hashlib
from django.core.management.base import BaseCommand
from django.db import connection, close_old_connections

SHEET_TO_REGION = {
    'גוש תחתון': {'region_id': 'gush-tachton', 'region_name': 'מעונות גוש תחתון'},
    'קנדה':      {'region_id': 'canada',        'region_name': 'מעונות קנדה'},
    'ברושים':    {'region_id': 'broshim',       'region_name': 'מעונות ברושים'},
    'מזרח':      {'region_id': 'mizrah',        'region_name': 'מעונות מזרח'},
    'גוש עליון': {'region_id': 'gush-elyon',   'region_name': 'מעונות גוש עליון'},
    'סגל זוטר':  {'region_id': 'segal-zutar',  'region_name': 'מעונות סגל זוטר'},
}


def reconnect():
    close_old_connections()
    connection.close()
    connection.ensure_connection()


def make_dorm_code(sheet_name, title):
    text = f"{sheet_name}-{title}".encode("utf-8")
    return int(hashlib.md5(text).hexdigest()[:6], 16) % 90000 + 10000


def get_capacity_from_title(title):
    """If title says 'שניים בחדר', every room has capacity 2."""
    if 'שניים בחדר' in str(title):
        return 2
    return None  # None means: count from rows


def get_category(apt_type_raw):
    t = str(apt_type_raw or '').strip()
    if 'בנות' in t or 'רווקה' in t:
        return 'female'
    return 'male'


def get_apartment_type(apt_type_raw):
    t = str(apt_type_raw or '').strip()
    if 'משפחות' in t or 'משפחה' in t:
        return 'family'
    if 'זוגות' in t or 'רווק בדירת' in t or 'רווקה בדירת' in t:
        return 'couple'
    return 'single'


def is_header_row(row):
    vals = {str(c).strip() for c in row if c is not None}
    return 'בניין' in vals and 'מזהה חדר' in vals


def is_title_row(row):
    if is_header_row(row):
        return False
    non_empty = [c for c in row if c is not None and str(c).strip() and str(c).strip() != 'nan']
    if len(non_empty) != 1:
        return False
    val = str(non_empty[0]).strip()
    try:
        float(val)
        return False
    except ValueError:
        pass
    if val == 'בניין':
        return False
    return True


def get_title_text(row):
    for c in row:
        if c is not None and str(c).strip() and str(c).strip() != 'nan':
            return str(c).strip()
    return ''


def parse_room_id(room_id):
    room_id = str(room_id or '').strip().rstrip('/')
    parts = room_id.split('/')
    if len(parts) >= 3 and parts[2]:
        return parts[1], parts[2]
    elif len(parts) >= 2 and parts[1]:
        return parts[1], '1'
    return None


class Command(BaseCommand):
    help = 'Import buildings. Optionally specify sheet name.'

    def add_arguments(self, parser):
        parser.add_argument('excel_path', type=str)
        parser.add_argument('sheet_name', nargs='?', default=None)

    def handle(self, *args, **options):
        import openpyxl
        from api.models import Region, DormType, Building, Apartment, Room, Bed

        excel_path = options['excel_path']
        target_sheet = options['sheet_name']

        self.stdout.write(f'📂 Reading: {excel_path}')
        wb = openpyxl.load_workbook(excel_path)

        total_buildings = total_apartments = total_rooms = total_beds = 0

        sheets_to_process = [target_sheet] if target_sheet else list(SHEET_TO_REGION.keys())

        for sheet_name in sheets_to_process:
            if sheet_name not in SHEET_TO_REGION:
                self.stdout.write(f'  ⚠️  Unknown sheet: {sheet_name}')
                continue
            if sheet_name not in wb.sheetnames:
                self.stdout.write(f'  ⚠️  Sheet not found: {sheet_name}')
                continue

            info = SHEET_TO_REGION[sheet_name]
            self.stdout.write(f'\n🏢 Processing sheet: {sheet_name}')

            reconnect()
            region, _ = Region.objects.get_or_create(
                id=info['region_id'],
                defaults={'name': info['region_name']}
            )

            ws = wb[sheet_name]
            all_rows = list(ws.iter_rows(values_only=True))

            sections = []
            current_title = info['region_name']
            current_col = {}
            current_rows = []

            for row in all_rows:
                if is_title_row(row):
                    if current_rows and current_col:
                        sections.append((current_title, current_col, current_rows))
                        current_rows = []
                    current_title = get_title_text(row)
                    current_col = {}
                elif is_header_row(row):
                    if current_rows and current_col:
                        sections.append((current_title, current_col, current_rows))
                        current_rows = []
                    current_col = {}
                    for j, h in enumerate(row):
                        h = str(h or '').strip()
                        if h == 'בניין': current_col['building'] = j
                        elif h == 'דירה': current_col['apt'] = j
                        elif h == 'חדר': current_col['room'] = j
                        elif h == 'מזהה חדר': current_col['room_id'] = j
                        elif h == 'סוג דירה': current_col['type'] = j
                        elif h == 'אוכלוסיה': current_col['population'] = j
                else:
                    if current_col and 'room_id' in current_col:
                        current_rows.append(row)

            if current_rows and current_col:
                sections.append((current_title, current_col, current_rows))

            self.stdout.write(f'  Found {len(sections)} sub-tables')

            for title, col, rows in sections:
                title_capacity = get_capacity_from_title(title)
                self.stdout.write(f'  📋 "{title}" → capacity_from_title={title_capacity}')

                reconnect()
                code = make_dorm_code(sheet_name, title)
                dorm_type, _ = DormType.objects.get_or_create(
                    name=title,
                    defaults={'region': region, 'code': code}
                )

                data = {}

                for row in rows:
                    bld_raw = row[col['building']] if 'building' in col and col['building'] < len(row) else None
                    if not bld_raw:
                        continue
                    try:
                        bld_num = int(float(str(bld_raw).strip()))
                    except (ValueError, TypeError):
                        continue

                    room_id_raw = row[col['room_id']] if 'room_id' in col and col['room_id'] < len(row) else None
                    room_id = str(room_id_raw or '').strip()
                    if not room_id or room_id == 'nan':
                        continue

                    parsed = parse_room_id(room_id)
                    if not parsed:
                        continue
                    apt_num, room_num = parsed

                    apt_type_raw = ''
                    if 'type' in col and col['type'] < len(row):
                        apt_type_raw = str(row[col['type']] or '').strip()

                    category = get_category(apt_type_raw)
                    apartment_type = get_apartment_type(apt_type_raw)

                    if bld_num not in data:
                        data[bld_num] = {}
                    if apt_num not in data[bld_num]:
                        data[bld_num][apt_num] = {
                            'rooms': {},
                            'category': category,
                            'apartment_type': apartment_type,
                        }

                    if room_num not in data[bld_num][apt_num]['rooms']:
                        data[bld_num][apt_num]['rooms'][room_num] = 0
                    data[bld_num][apt_num]['rooms'][room_num] += 1

                for building_num, apts in data.items():
                    reconnect()
                    building, created = Building.objects.get_or_create(
                        number=building_num,
                        dorm_type=dorm_type,
                        defaults={'is_active': True}
                    )
                    if created:
                        total_buildings += 1

                    for apt_num, apt_data in apts.items():
                        rooms_data = apt_data['rooms']
                        room_count = len(rooms_data)

                        # Use title capacity if specified, else sum from rows
                        if title_capacity is not None:
                            apartment_capacity = room_count * title_capacity
                        else:
                            apartment_capacity = sum(rooms_data.values())

                        apartment, created = Apartment.objects.get_or_create(
                            building=building,
                            number=str(apt_num),
                            defaults={
                                'category': apt_data['category'],
                                'apartment_type': apt_data['apartment_type'],
                                'room_count': room_count,
                                'apartment_capacity': apartment_capacity,
                                'is_active': True,
                                'inactive_reason': '',
                            }
                        )
                        if not created:
                            apartment.apartment_capacity = apartment_capacity
                            apartment.room_count = room_count
                            apartment.save()
                        else:
                            total_apartments += 1

                        for room_num, row_count in rooms_data.items():
                            # Final capacity: title overrides row count
                            capacity = title_capacity if title_capacity is not None else row_count

                            room, created = Room.objects.get_or_create(
                                apartment=apartment,
                                name=str(room_num),
                                defaults={'capacity': capacity, 'is_active': True}
                            )
                            if not created:
                                if room.capacity != capacity:
                                    room.capacity = capacity
                                    room.save()
                            else:
                                total_rooms += 1

                            reconnect()
                            existing_beds = set(room.beds.values_list('label', flat=True))
                            new_beds = {f'Bed {i+1}' for i in range(capacity)}

                            for old_label in existing_beds - new_beds:
                                room.beds.filter(label=old_label).delete()

                            for bed_label in new_beds - existing_beds:
                                Bed.objects.create(room=room, label=bed_label)
                                total_beds += 1

                self.stdout.write(f'    ✅ {len(data)} buildings')

        self.stdout.write(self.style.SUCCESS(
            f'\n🎉 Import Complete!'
            f'\n   🏢 Buildings:  {total_buildings}'
            f'\n   🏠 Apartments: {total_apartments}'
            f'\n   🚪 Rooms:      {total_rooms}'
            f'\n   🛏️  Beds:       {total_beds}'
        ))