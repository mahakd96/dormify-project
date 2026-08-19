from io import BytesIO

import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from .models import BedAssignment, Bed, Student


# ---------------------------------------------------------------------------
# Formatting constants
# ---------------------------------------------------------------------------

_HEADER_FILL  = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
_HEADER_FONT  = Font(bold=True, color='1F3864', size=11)
_HEADER_ALIGN = Alignment(horizontal='right', vertical='center', wrap_text=False)
_CELL_ALIGN   = Alignment(horizontal='right', vertical='center', wrap_text=False)

# Row fill colors
_FILL_UNASSIGNED = PatternFill(start_color='FFF3E0', end_color='FFF3E0', fill_type='solid')
_FILL_EXCEPTION  = PatternFill(start_color='FFEBEE', end_color='FFEBEE', fill_type='solid')
_FILL_LEAVING    = PatternFill(start_color='ECEFF1', end_color='ECEFF1', fill_type='solid')
_FILL_NEW        = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
_FILL_CONTINUING = PatternFill(start_color='E3F2FD', end_color='E3F2FD', fill_type='solid')
_FILL_TRANSFER   = PatternFill(start_color='FFF8E1', end_color='FFF8E1', fill_type='solid')
_FILL_PRIORITY   = PatternFill(start_color='EDE7F6', end_color='EDE7F6', fill_type='solid')
_FILL_OCC_LOW    = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
_FILL_OCC_MED    = PatternFill(start_color='FFF8E1', end_color='FFF8E1', fill_type='solid')
_FILL_OCC_HIGH   = PatternFill(start_color='FFEBEE', end_color='FFEBEE', fill_type='solid')
_FILL_DUP_ID     = PatternFill(start_color='FCE4EC', end_color='FCE4EC', fill_type='solid')
_FILL_MISSING    = PatternFill(start_color='FBE9E7', end_color='FBE9E7', fill_type='solid')

# Tab colors (hex, no #)
_TAB_FULL        = '1B5E20'
_TAB_STUDENT     = '1565C0'
_TAB_CAPACITY    = '6A1B9A'
_TAB_REVIEW      = 'B71C1C'
_TAB_GRAY        = '607D8B'
_TAB_ORANGE      = 'E65100'
_TAB_GREEN       = '2E7D32'
_TAB_LIGHT_BLUE  = '0277BD'
_TAB_AMBER       = 'F57F17'
_TAB_TEAL        = '00695C'
_TAB_RED         = 'C62828'
_TAB_DARK_ORANGE = 'BF360C'


# ---------------------------------------------------------------------------
# Core helpers
# ---------------------------------------------------------------------------

def _safe(value, fallback=''):
    if value is None:
        return fallback
    return str(value)


def _to_df(rows):
    if not rows:
        return pd.DataFrame([{'הערה': 'אין נתונים'}])
    return pd.DataFrame(rows)


def _expl_df(lines):
    return pd.DataFrame([{'פרטים': line} for line in lines])


def _set_tab_color(ws, hex_color):
    ws.sheet_properties.tabColor = hex_color


def _format_worksheet(ws):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = 'A2'

    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font      = _HEADER_FONT
        cell.fill      = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _CELL_ALIGN

    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max(
            (len(str(cell.value)) if cell.value is not None else 0)
            for cell in col_cells
        )
        ws.column_dimensions[col_letter].width = min(max(round(max_len * 1.3) + 2, 10), 45)


def _reapply_headers(book):
    """Re-apply header styling after row fills (fills may overwrite headers)."""
    for ws in book.worksheets:
        for cell in ws[1]:
            cell.font      = _HEADER_FONT
            cell.fill      = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN


def _apply_uniform_row_fill(ws, fill):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.fill = fill


def _highlight_priority_rows(ws):
    header = {cell.value: cell.column for cell in ws[1]}
    col = header.get('עדיפות')
    if col is None:
        return
    for row in ws.iter_rows(min_row=2):
        if str(row[col - 1].value or '').strip() == 'כן':
            for cell in row:
                cell.fill = _FILL_PRIORITY


def _color_occupancy_rows(ws):
    header = {cell.value: cell.column for cell in ws[1]}
    col = header.get('אחוז תפוסה')
    if col is None:
        return
    for row in ws.iter_rows(min_row=2):
        try:
            pct = float(str(row[col - 1].value or '').replace('%', '').strip())
        except (ValueError, TypeError):
            continue
        fill = _FILL_OCC_HIGH if pct >= 90 else (_FILL_OCC_MED if pct >= 70 else _FILL_OCC_LOW)
        for cell in row:
            cell.fill = fill


# ---------------------------------------------------------------------------
# Shared row builders
# ---------------------------------------------------------------------------

def _location_from_assignment(assignment):
    if not assignment:
        return {'אזור': '', 'סוג מעון': '', 'בניין': '', 'דירה': '', 'חדר': '', 'מיטה': ''}
    bed      = assignment.bed
    room     = bed.room
    apt      = room.apartment
    building = apt.building
    dorm_type = building.dorm_type
    region   = dorm_type.region if dorm_type else None
    return {
        'אזור':     _safe(region.name if region else ''),
        'סוג מעון': _safe(dorm_type.name if dorm_type else ''),
        'בניין':    _safe(building.number),
        'דירה':     _safe(apt.number),
        'חדר':      _safe(room.name),
        'מיטה':     _safe(bed.label),
    }


def _student_base(student):
    return {
        'תעודת זהות': _safe(student.student_id),
        'שם מלא':     _safe(student.full_name),
        'מגדר':       student.get_gender_display() if student.gender else '',
        'דת':         student.get_requested_religion_display() if student.requested_religion else '',
        'קטגוריה':    student.get_category_display() if student.category else '',
        'עדיפות':     'כן' if student.is_priority else 'לא',
    }


def _student_issues(student, seen_ids):
    """Return list of data-quality issues for a student. Mutates seen_ids."""
    issues = []
    sid = student.student_id
    if not sid:
        issues.append('חסרה תעודת זהות')
    elif sid in seen_ids:
        issues.append('תעודת זהות כפולה')
    else:
        seen_ids[sid] = student.full_name

    if not student.first_name or not student.last_name:
        issues.append('חסר שם')
    if not student.gender:
        issues.append('חסר מגדר')
    if not student.housing_type and not student.accepted_dorm_type_id:
        issues.append('חסר סוג דיור / מעון')
    return issues


# ---------------------------------------------------------------------------
# Data loading with optional region filter
# ---------------------------------------------------------------------------

def _load_data(region_id=None, include_students=True):
    """
    Load beds + occupied_bed_ids always; students + active assignments
    only when include_students=True (the default - preserves every
    pre-existing caller's exact prior behavior, including the dict's key
    set, which is always the same regardless of include_students).

    generate_capacity_report() passes include_students=False: it renders
    bed/room/building/apartment occupancy only and never reads
    all_students/active_assignments/assignment_map/assigned_pks/the
    category-bucket lists - skipping that fetch removes an entirely
    unused Student+BedAssignment query pair (plus their 5 Python-side
    category list comprehensions) from every capacity-report generation,
    with zero effect on any other report (dormify_report,
    student_actions_report, manual_review_report all keep the
    include_students=True default, unchanged).
    (project-quality/performance/PERFORMANCE_FINAL_REPORT.md, G1-20.)
    """
    active_assignments  = []
    assignment_map      = {}
    assigned_pks        = set()
    all_students        = []
    new_students        = []
    continuing_students  = []
    transfer_students    = []
    leaving_students     = []
    unassigned_students  = []

    if include_students:
        assignment_qs = (
            BedAssignment.objects.filter(status=BedAssignment.Status.ACTIVE)
            .select_related(
                'student__accepted_dorm_type__region',
                'bed__room__apartment__building__dorm_type__region',
            )
        )
        student_qs = Student.objects.select_related('accepted_dorm_type__region')

        if region_id:
            assignment_qs = assignment_qs.filter(
                bed__room__apartment__building__dorm_type__region_id=region_id
            )
            student_qs = student_qs.filter(accepted_dorm_type__region_id=region_id)

        active_assignments = list(assignment_qs)
        assignment_map = {a.student_id: a for a in active_assignments}
        assigned_pks   = set(assignment_map.keys())
        all_students   = list(student_qs.all())

        cat = Student.StudentCategory
        new_students        = [s for s in all_students if s.category == cat.NEW]
        continuing_students = [s for s in all_students if s.category == cat.CONTINUING]
        transfer_students   = [s for s in all_students if s.category == cat.TRANSFER]
        leaving_students    = [s for s in all_students if s.category == cat.LEAVING]
        unassigned_students = [s for s in all_students if s.pk not in assigned_pks]

    bed_qs = Bed.objects.select_related('room__apartment__building__dorm_type__region')
    if region_id:
        bed_qs = bed_qs.filter(room__apartment__building__dorm_type__region_id=region_id)
    all_beds = list(bed_qs.all())

    occ_qs = BedAssignment.objects.filter(status=BedAssignment.Status.ACTIVE)
    if region_id:
        occ_qs = occ_qs.filter(
            bed__room__apartment__building__dorm_type__region_id=region_id
        )
    occupied_bed_ids = set(occ_qs.values_list('bed_id', flat=True))

    return {
        'active_assignments':   active_assignments,
        'assignment_map':       assignment_map,
        'assigned_pks':         assigned_pks,
        'all_students':         all_students,
        'all_beds':             all_beds,
        'occupied_bed_ids':     occupied_bed_ids,
        'new_students':         new_students,
        'continuing_students':  continuing_students,
        'transfer_students':    transfer_students,
        'leaving_students':     leaving_students,
        'unassigned_students':  unassigned_students,
    }


# Keep backward-compat alias
def _load_common_data():
    return _load_data(region_id=None)


# ---------------------------------------------------------------------------
# Full Dormify Report builders (unchanged from original)
# ---------------------------------------------------------------------------

def _build_summary(all_students, active_assignments, all_beds, occupied_bed_ids):
    total_students   = len(all_students)
    assigned_pks     = {a.student_id for a in active_assignments}
    assigned_count   = len(assigned_pks)
    unassigned_count = total_students - assigned_count

    total_beds    = len(all_beds)
    occupied_beds = len(occupied_bed_ids)
    occ_pct       = round(occupied_beds / total_beds * 100, 1) if total_beds else 0

    gender_counts   = {}
    category_counts = {}
    region_counts   = {}

    for s in all_students:
        g = s.get_gender_display() if s.gender else 'לא צוין'
        gender_counts[g] = gender_counts.get(g, 0) + 1
        c = s.get_category_display() if s.category else 'לא צוין'
        category_counts[c] = category_counts.get(c, 0) + 1

    for a in active_assignments:
        try:
            name = a.bed.room.apartment.building.dorm_type.region.name
        except AttributeError:
            name = 'לא ידוע'
        region_counts[name] = region_counts.get(name, 0) + 1

    rows = [
        {'נושא': 'סה"כ סטודנטים',      'ערך': total_students},
        {'נושא': 'סטודנטים משובצים',    'ערך': assigned_count},
        {'נושא': 'סטודנטים לא משובצים', 'ערך': unassigned_count},
        {'נושא': '', 'ערך': ''},
        {'נושא': 'לפי קטגוריה', 'ערך': ''},
    ]
    for lbl, cnt in sorted(category_counts.items()):
        rows.append({'נושא': f'  {lbl}', 'ערך': cnt})

    rows += [{'נושא': '', 'ערך': ''}, {'נושא': 'לפי מגדר', 'ערך': ''}]
    for lbl, cnt in sorted(gender_counts.items()):
        rows.append({'נושא': f'  {lbl}', 'ערך': cnt})

    rows += [{'נושא': '', 'ערך': ''}, {'נושא': 'לפי אזור (משובצים)', 'ערך': ''}]
    for lbl, cnt in sorted(region_counts.items()):
        rows.append({'נושא': f'  {lbl}', 'ערך': cnt})

    rows += [
        {'נושא': '', 'ערך': ''},
        {'נושא': 'סה"כ מיטות',   'ערך': total_beds},
        {'נושא': 'מיטות תפוסות', 'ערך': occupied_beds},
        {'נושא': 'מיטות פנויות', 'ערך': total_beds - occupied_beds},
        {'נושא': 'אחוז תפוסה',   'ערך': f'{occ_pct}%'},
    ]
    return rows


def _build_assignment_rows(active_assignments):
    rows = []
    for a in active_assignments:
        row = {**_student_base(a.student), **_location_from_assignment(a), 'סטטוס שיבוץ': 'משובץ'}
        rows.append(row)
    return rows


def _build_unassigned_rows(unassigned_students):
    rows = []
    for student in unassigned_students:
        dorm_type = student.accepted_dorm_type
        region    = dorm_type.region if dorm_type else None
        rows.append({
            **_student_base(student),
            'סוג מעון מבוקש': _safe(dorm_type.name if dorm_type else ''),
            'אזור':           _safe(region.name if region else ''),
            'הערות':          _safe(student.priority_reason),
        })
    return rows


def _build_occupancy_rows(all_beds, occupied_bed_ids):
    room_data = {}
    for bed in all_beds:
        room    = bed.room
        rid     = room.pk
        if rid not in room_data:
            apt       = room.apartment
            building  = apt.building
            dorm_type = building.dorm_type
            region    = dorm_type.region if dorm_type else None
            room_data[rid] = {
                'אזור':     _safe(region.name if region else ''),
                'סוג מעון': _safe(dorm_type.name if dorm_type else ''),
                'בניין':    _safe(building.number),
                'דירה':     _safe(apt.number),
                'חדר':      _safe(room.name),
                'total': 0, 'occupied': 0,
            }
        room_data[rid]['total'] += 1
        if bed.pk in occupied_bed_ids:
            room_data[rid]['occupied'] += 1

    rows = []
    for rd in sorted(room_data.values(), key=lambda x: (x['אזור'], x['בניין'], x['דירה'], x['חדר'])):
        t, o = rd['total'], rd['occupied']
        rows.append({
            'אזור':          rd['אזור'],
            'סוג מעון':      rd['סוג מעון'],
            'בניין':         rd['בניין'],
            'דירה':          rd['דירה'],
            'חדר':           rd['חדר'],
            'סה"כ מיטות':   t,
            'מיטות תפוסות': o,
            'מיטות פנויות': t - o,
            'אחוז תפוסה':   f'{round(o / t * 100, 1) if t else 0}%',
        })
    return rows


def _build_category_rows(students, assignment_map):
    rows = []
    for student in students:
        assignment = assignment_map.get(student.pk)
        rows.append({**_student_base(student), **_location_from_assignment(assignment)})
    return rows


def _build_exceptions_rows(all_students, assignment_map):
    rows     = []
    seen_ids = {}
    for student in all_students:
        issues = _student_issues(student, seen_ids)
        if issues:
            rows.append({**_student_base(student), 'סיבות חריגה': ' | '.join(issues)})
    return rows


# ---------------------------------------------------------------------------
# Student Actions Report builders
# ---------------------------------------------------------------------------

def _classify_students(all_students, assignment_map):
    """Route each student to exactly one action bucket. No duplication."""
    cat = Student.StudentCategory
    buckets = {
        'נכנסים חדשים': [],
        'ממשיכים':       [],
        'מעברים':        [],
        'עוזבים':        [],
        'לא שובצו':      [],
    }
    for student in all_students:
        assignment = assignment_map.get(student.pk)
        if not assignment:
            buckets['לא שובצו'].append(student)
        elif student.category == cat.LEAVING:
            buckets['עוזבים'].append(student)
        elif student.category == cat.NEW:
            buckets['נכנסים חדשים'].append(student)
        elif student.category == cat.CONTINUING:
            buckets['ממשיכים'].append(student)
        elif student.category == cat.TRANSFER:
            buckets['מעברים'].append(student)
        else:
            # Assigned but unknown category → treat as continuing
            buckets['ממשיכים'].append(student)
    return buckets


def _build_action_rows(students, assignment_map):
    """Build uniform rows for any action sheet. Columns consistent across all sheets."""
    rows = []
    for student in students:
        assignment = assignment_map.get(student.pk)
        if assignment:
            loc    = _location_from_assignment(assignment)
            status = 'משובץ'
        else:
            dorm_type = student.accepted_dorm_type
            region    = dorm_type.region if dorm_type else None
            loc = {
                'אזור':     _safe(region.name if region else ''),
                'סוג מעון': _safe(dorm_type.name if dorm_type else ''),
                'בניין': '', 'דירה': '', 'חדר': '', 'מיטה': '',
            }
            status = 'לא שובץ'

        rows.append({
            **_student_base(student),
            **loc,
            'סטטוס שיבוץ': status,
            'הערות':        _safe(student.priority_reason),
        })
    return rows


def _build_actions_summary_rows(all_students, classified, active_assignments, region_name=None):
    assigned_pks = {a.student_id for a in active_assignments}
    total    = len(all_students)
    assigned = len(assigned_pks)

    gender_counts = {}
    for s in all_students:
        g = s.get_gender_display() if s.gender else 'לא צוין'
        gender_counts[g] = gender_counts.get(g, 0) + 1

    rows = [
        {'נושא': 'אזור נבחר',              'ערך': region_name or 'כל האזורים'},
        {'נושא': '', 'ערך': ''},
        {'נושא': 'סה"כ סטודנטים',          'ערך': total},
        {'נושא': 'משובצים',                 'ערך': assigned},
        {'נושא': 'לא משובצים',              'ערך': total - assigned},
        {'נושא': '', 'ערך': ''},
        {'נושא': 'לפי פעולה / קטגוריה',   'ערך': ''},
        {'נושא': '  נכנסים חדשים',          'ערך': len(classified['נכנסים חדשים'])},
        {'נושא': '  ממשיכים',               'ערך': len(classified['ממשיכים'])},
        {'נושא': '  מעברים',                'ערך': len(classified['מעברים'])},
        {'נושא': '  עוזבים',                'ערך': len(classified['עוזבים'])},
        {'נושא': '  לא שובצו',              'ערך': len(classified['לא שובצו'])},
        {'נושא': '', 'ערך': ''},
        {'נושא': 'לפי מגדר',               'ערך': ''},
    ]
    for g, c in sorted(gender_counts.items()):
        rows.append({'נושא': f'  {g}', 'ערך': c})

    rows += [
        {'נושא': '', 'ערך': ''},
        {'נושא': 'בעלי עדיפות', 'ערך': sum(1 for s in all_students if s.is_priority)},
    ]
    return rows


# ---------------------------------------------------------------------------
# Capacity Report builders
# ---------------------------------------------------------------------------

def _build_capacity_summary_rows(all_beds, occupied_bed_ids, region_name=None):
    regions   = set()
    buildings = set()
    apartments = set()
    rooms     = set()

    for bed in all_beds:
        room      = bed.room
        apt       = room.apartment
        building  = apt.building
        dorm_type = building.dorm_type
        region    = dorm_type.region if dorm_type else None
        rooms.add(room.pk)
        apartments.add(apt.pk)
        buildings.add(building.pk)
        if region:
            regions.add(region.pk)

    total    = len(all_beds)
    occupied = len(occupied_bed_ids)
    pct      = round(occupied / total * 100, 1) if total else 0

    return [
        {'נושא': 'אזור נבחר',    'ערך': region_name or 'כל האזורים'},
        {'נושא': '', 'ערך': ''},
        {'נושא': 'סה"כ מיטות',   'ערך': total},
        {'נושא': 'מיטות תפוסות', 'ערך': occupied},
        {'נושא': 'מיטות פנויות', 'ערך': total - occupied},
        {'נושא': 'אחוז תפוסה',   'ערך': f'{pct}%'},
        {'נושא': '', 'ערך': ''},
        {'נושא': 'מספר אזורים',  'ערך': len(regions)},
        {'נושא': 'מספר בניינים', 'ערך': len(buildings)},
        {'נושא': 'מספר דירות',   'ערך': len(apartments)},
        {'נושא': 'מספר חדרים',   'ערך': len(rooms)},
    ]


def _build_occupancy_by_region(all_beds, occupied_bed_ids):
    data = {}
    for bed in all_beds:
        dorm_type = bed.room.apartment.building.dorm_type
        region    = dorm_type.region if dorm_type else None
        name      = region.name if region else 'לא ידוע'
        if name not in data:
            data[name] = {'total': 0, 'occupied': 0}
        data[name]['total'] += 1
        if bed.pk in occupied_bed_ids:
            data[name]['occupied'] += 1

    rows = []
    for name, d in sorted(data.items()):
        t, o = d['total'], d['occupied']
        rows.append({
            'אזור':          name,
            'סה"כ מיטות':   t,
            'מיטות תפוסות': o,
            'מיטות פנויות': t - o,
            'אחוז תפוסה':   f'{round(o / t * 100, 1) if t else 0}%',
        })
    return rows


def _build_occupancy_by_building(all_beds, occupied_bed_ids):
    data = {}
    for bed in all_beds:
        room      = bed.room
        apt       = room.apartment
        building  = apt.building
        dorm_type = building.dorm_type
        region    = dorm_type.region if dorm_type else None
        key       = building.pk
        if key not in data:
            data[key] = {
                'אזור':     _safe(region.name if region else ''),
                'סוג מעון': _safe(dorm_type.name if dorm_type else ''),
                'בניין':    _safe(building.number),
                '_sort':    (_safe(region.name if region else ''), _safe(building.number)),
                'total': 0, 'occupied': 0,
            }
        data[key]['total'] += 1
        if bed.pk in occupied_bed_ids:
            data[key]['occupied'] += 1

    rows = []
    for d in sorted(data.values(), key=lambda x: x['_sort']):
        t, o = d['total'], d['occupied']
        rows.append({
            'אזור':          d['אזור'],
            'סוג מעון':      d['סוג מעון'],
            'בניין':         d['בניין'],
            'סה"כ מיטות':   t,
            'מיטות תפוסות': o,
            'מיטות פנויות': t - o,
            'אחוז תפוסה':   f'{round(o / t * 100, 1) if t else 0}%',
        })
    return rows


def _build_occupancy_by_apartment(all_beds, occupied_bed_ids):
    data = {}
    for bed in all_beds:
        room      = bed.room
        apt       = room.apartment
        building  = apt.building
        dorm_type = building.dorm_type
        region    = dorm_type.region if dorm_type else None
        key       = apt.pk
        if key not in data:
            data[key] = {
                'אזור':  _safe(region.name if region else ''),
                'בניין': _safe(building.number),
                'דירה':  _safe(apt.number),
                '_sort': (_safe(region.name if region else ''), _safe(building.number), _safe(apt.number)),
                'total': 0, 'occupied': 0,
            }
        data[key]['total'] += 1
        if bed.pk in occupied_bed_ids:
            data[key]['occupied'] += 1

    rows = []
    for d in sorted(data.values(), key=lambda x: x['_sort']):
        t, o = d['total'], d['occupied']
        rows.append({
            'אזור':          d['אזור'],
            'בניין':         d['בניין'],
            'דירה':          d['דירה'],
            'סה"כ מיטות':   t,
            'מיטות תפוסות': o,
            'מיטות פנויות': t - o,
            'אחוז תפוסה':   f'{round(o / t * 100, 1) if t else 0}%',
        })
    return rows


# ---------------------------------------------------------------------------
# Manual Review Report builders
# ---------------------------------------------------------------------------

def _compute_student_issues(all_students):
    """
    Computes _student_issues() for every student ONCE, in a single pass
    with one shared seen_ids dict, instead of the 3 separate call sites
    below (_build_exceptions_summary_rows, _build_all_exceptions_rows,
    _build_unassigned_with_issues) each independently re-running it with
    their own fresh seen_ids over the same all_students list.

    Safe: all three original call sites iterated the exact same
    already-materialized `all_students` list (fetched once by
    _load_data(), never re-queried or re-ordered between builder calls)
    in the exact same order, each with a FRESH seen_ids - so a given
    student's issues list (including whether THAT occurrence is flagged
    'תעודת זהות כפולה') was always identical across all three calls to
    begin with; this just computes that one identical result once and
    reuses it, instead of recomputing it 3 times.
    (project-quality/performance/PERFORMANCE_FINAL_REPORT.md, G1-19.)

    Returns a list of (student, issues) pairs, same order as all_students.
    """
    seen_ids = {}
    return [(student, _student_issues(student, seen_ids)) for student in all_students]


def _build_exceptions_summary_rows(student_issues, assignment_map, region_name=None):
    missing_id          = 0
    dup_id              = 0
    missing_name        = 0
    missing_gender      = 0
    missing_dorm        = 0
    total_exc           = 0
    unassigned_w_issues = 0

    for student, issues in student_issues:
        if issues:
            total_exc += 1
            if student.pk not in assignment_map:
                unassigned_w_issues += 1

        if not student.student_id:
            missing_id += 1
        if 'תעודת זהות כפולה' in issues:
            dup_id += 1
        if not student.first_name or not student.last_name:
            missing_name += 1
        if not student.gender:
            missing_gender += 1
        if not student.housing_type and not student.accepted_dorm_type_id:
            missing_dorm += 1

    return [
        {'נושא': 'אזור נבחר',               'ערך': region_name or 'כל האזורים'},
        {'נושא': '', 'ערך': ''},
        {'נושא': 'סה"כ רשומות עם חריגות',  'ערך': total_exc},
        {'נושא': '', 'ערך': ''},
        {'נושא': 'חסרה תעודת זהות',         'ערך': missing_id},
        {'נושא': 'תעודת זהות כפולה',        'ערך': dup_id},
        {'נושא': 'חסר שם',                  'ערך': missing_name},
        {'נושא': 'חסר מגדר',                'ערך': missing_gender},
        {'נושא': 'חסר סוג דיור / מעון',     'ערך': missing_dorm},
        {'נושא': '', 'ערך': ''},
        {'נושא': 'לא משובצים עם בעיה',     'ערך': unassigned_w_issues},
    ]


def _build_all_exceptions_rows(student_issues, assignment_map):
    rows = []
    for student, issues in student_issues:
        if issues:
            dorm_type = student.accepted_dorm_type
            region    = dorm_type.region if dorm_type else None
            rows.append({
                **_student_base(student),
                'אזור':         _safe(region.name if region else ''),
                'סיבות חריגה': ' | '.join(issues),
                'סטטוס שיבוץ': 'משובץ' if student.pk in assignment_map else 'לא שובץ',
                'הערות':        _safe(student.priority_reason),
            })
    return rows


def _build_duplicate_id_rows(all_students):
    groups = {}
    for student in all_students:
        if student.student_id:
            groups.setdefault(student.student_id, []).append(student)

    rows = []
    for sid, students in sorted(groups.items()):
        if len(students) > 1:
            for s in students:
                rows.append({
                    **_student_base(s),
                    'הערה': f'כפילות – {len(students)} רשומות',
                })
    return rows


def _build_missing_data_rows(all_students):
    rows = []
    for student in all_students:
        issues = []
        if not student.student_id:
            issues.append('חסרה תעודת זהות')
        if not student.first_name or not student.last_name:
            issues.append('חסר שם')
        if not student.gender:
            issues.append('חסר מגדר')
        if not student.housing_type and not student.accepted_dorm_type_id:
            issues.append('חסר סוג דיור')
        if issues:
            rows.append({**_student_base(student), 'נתונים חסרים': ' | '.join(issues)})
    return rows


def _build_unassigned_with_issues(student_issues, assignment_map):
    """Only unassigned students who ALSO have a data-quality issue."""
    rows = []
    for student, issues in student_issues:
        if issues and student.pk not in assignment_map:
            dorm_type = student.accepted_dorm_type
            region    = dorm_type.region if dorm_type else None
            rows.append({
                **_student_base(student),
                'אזור':         _safe(region.name if region else ''),
                'סיבות חריגה': ' | '.join(issues),
                'סטטוס שיבוץ': 'לא שובץ',
                'הערות':        _safe(student.priority_reason),
            })
    return rows


# ---------------------------------------------------------------------------
# Report 1: Full Dormify Report (backward-compat, unchanged logic)
# ---------------------------------------------------------------------------

def generate_dormify_report():
    d = _load_data()

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        _to_df(
            _build_summary(d['all_students'], d['active_assignments'], d['all_beds'], d['occupied_bed_ids'])
        ).to_excel(writer, sheet_name='סיכום', index=False)

        _to_df(_build_assignment_rows(d['active_assignments'])
               ).to_excel(writer, sheet_name='דוח שיבוץ', index=False)

        _to_df(_build_unassigned_rows(d['unassigned_students'])
               ).to_excel(writer, sheet_name='לא שובצו', index=False)

        _to_df(_build_occupancy_rows(d['all_beds'], d['occupied_bed_ids'])
               ).to_excel(writer, sheet_name='תפוסה', index=False)

        _to_df(_build_category_rows(d['new_students'], d['assignment_map'])
               ).to_excel(writer, sheet_name='נכנסים חדשים', index=False)

        _to_df(_build_category_rows(d['continuing_students'], d['assignment_map'])
               ).to_excel(writer, sheet_name='ממשיכים', index=False)

        _to_df(_build_category_rows(d['transfer_students'], d['assignment_map'])
               ).to_excel(writer, sheet_name='מעברים', index=False)

        _to_df(_build_category_rows(d['leaving_students'], d['assignment_map'])
               ).to_excel(writer, sheet_name='עוזבים', index=False)

        _to_df(_build_exceptions_rows(d['all_students'], d['assignment_map'])
               ).to_excel(writer, sheet_name='חריגים לבדיקה', index=False)

        for ws in writer.book.worksheets:
            _format_worksheet(ws)
            _set_tab_color(ws, _TAB_FULL)

    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# Report 2: Student Actions Report  (replaces Student Allocation Report)
# ---------------------------------------------------------------------------

def generate_student_actions_report(region_id=None, region_name=None):
    """
    Students grouped by office action: new / continuing / transfer / leaving / unassigned.
    Each student appears in exactly ONE sheet — no duplication.
    Optionally filtered to a single region via region_id.
    """
    d          = _load_data(region_id=region_id)
    classified = _classify_students(d['all_students'], d['assignment_map'])
    region_lbl = region_name or ('כל האזורים' if not region_id else region_id)

    expl_lines = [
        'דוח פעולות סטודנטים – Dormify',
        '',
        f'אזור נבחר: {region_lbl}',
        '',
        'הדוח מציג את הסטודנטים לפי הפעולה הנדרשת על ידי המשרד.',
        'כל סטודנט מופיע בגיליון אחד בלבד – אין כפילויות.',
        '',
        'לתוצאות מעודכנות: העלה נתונים עדכניים והרץ שיבוץ לפני ההורדה.',
        '',
        'גיליון "נכנסים חדשים" – סטודנטים חדשים עם שיבוץ פעיל.',
        'גיליון "ממשיכים"       – סטודנטים ממשיכים עם שיבוץ פעיל.',
        'גיליון "מעברים"        – סטודנטים בתהליך מעבר עם שיבוץ פעיל.',
        'גיליון "עוזבים"        – סטודנטים עם שיבוץ פעיל שמסומנים כעוזבים.',
        'גיליון "לא שובצו"      – כל הסטודנטים ללא שיבוץ פעיל.',
    ]

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        _expl_df(expl_lines).to_excel(writer, sheet_name='הסבר', index=False)

        _to_df(_build_actions_summary_rows(
            d['all_students'], classified, d['active_assignments'], region_lbl
        )).to_excel(writer, sheet_name='תמצית סטודנטים', index=False)

        _to_df(_build_action_rows(classified['נכנסים חדשים'], d['assignment_map'])
               ).to_excel(writer, sheet_name='נכנסים חדשים', index=False)

        _to_df(_build_action_rows(classified['ממשיכים'], d['assignment_map'])
               ).to_excel(writer, sheet_name='ממשיכים', index=False)

        _to_df(_build_action_rows(classified['מעברים'], d['assignment_map'])
               ).to_excel(writer, sheet_name='מעברים', index=False)

        _to_df(_build_action_rows(classified['עוזבים'], d['assignment_map'])
               ).to_excel(writer, sheet_name='עוזבים', index=False)

        _to_df(_build_action_rows(classified['לא שובצו'], d['assignment_map'])
               ).to_excel(writer, sheet_name='לא שובצו', index=False)

        book = writer.book
        tab_colors = {
            'הסבר':           _TAB_GRAY,
            'תמצית סטודנטים': _TAB_STUDENT,
            'נכנסים חדשים':   _TAB_GREEN,
            'ממשיכים':         _TAB_LIGHT_BLUE,
            'מעברים':          _TAB_AMBER,
            'עוזבים':          _TAB_GRAY,
            'לא שובצו':        _TAB_ORANGE,
        }
        for ws in book.worksheets:
            _format_worksheet(ws)
            _set_tab_color(ws, tab_colors.get(ws.title, _TAB_STUDENT))

        _apply_uniform_row_fill(book['נכנסים חדשים'], _FILL_NEW)
        _apply_uniform_row_fill(book['ממשיכים'],       _FILL_CONTINUING)
        _apply_uniform_row_fill(book['מעברים'],        _FILL_TRANSFER)
        _apply_uniform_row_fill(book['עוזבים'],        _FILL_LEAVING)
        _apply_uniform_row_fill(book['לא שובצו'],      _FILL_UNASSIGNED)

        # Priority rows override category color
        for sheet_name in ['נכנסים חדשים', 'ממשיכים', 'מעברים', 'עוזבים', 'לא שובצו']:
            _highlight_priority_rows(book[sheet_name])

        _reapply_headers(book)

    buffer.seek(0)
    return buffer


# Keep old name as alias for backward compatibility
def generate_student_allocation_report():
    return generate_student_actions_report()


# ---------------------------------------------------------------------------
# Report 3: Capacity / Occupancy Report
# ---------------------------------------------------------------------------

def generate_capacity_report(region_id=None, region_name=None):
    # include_students=False: this report only reads d['all_beds'] and
    # d['occupied_bed_ids'] below - it never touches d['all_students'] or
    # any student/assignment-derived key, so the Student+BedAssignment
    # fetch _load_data() would otherwise do is entirely wasted work here.
    # (G1-20.)
    d          = _load_data(region_id=region_id, include_students=False)
    region_lbl = region_name or ('כל האזורים' if not region_id else region_id)

    expl_lines = [
        'דוח תפוסה ומיטות – Dormify',
        '',
        f'אזור נבחר: {region_lbl}',
        '',
        'הדוח מסכם את תפוסת המעונות ואינו כולל פרטי סטודנטים.',
        'הגיליונות הראשונים מציגים סיכומים; הגיליון האחרון מפרט ברמת חדר.',
        '',
        'ירוק  – תפוסה נמוכה (מתחת ל-70%)',
        'צהוב  – תפוסה בינונית (70%–89%)',
        'אדום  – תפוסה גבוהה (90% ומעלה)',
    ]

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        _expl_df(expl_lines).to_excel(writer, sheet_name='הסבר', index=False)

        _to_df(_build_capacity_summary_rows(d['all_beds'], d['occupied_bed_ids'], region_lbl)
               ).to_excel(writer, sheet_name='תמצית תפוסה', index=False)

        _to_df(_build_occupancy_by_region(d['all_beds'], d['occupied_bed_ids'])
               ).to_excel(writer, sheet_name='תפוסה לפי אזור', index=False)

        _to_df(_build_occupancy_by_building(d['all_beds'], d['occupied_bed_ids'])
               ).to_excel(writer, sheet_name='תפוסה לפי בניין', index=False)

        _to_df(_build_occupancy_by_apartment(d['all_beds'], d['occupied_bed_ids'])
               ).to_excel(writer, sheet_name='תפוסה לפי דירה', index=False)

        _to_df(_build_occupancy_rows(d['all_beds'], d['occupied_bed_ids'])
               ).to_excel(writer, sheet_name='תפוסה לפי חדרים', index=False)

        book = writer.book
        tab_colors = {
            'הסבר':            _TAB_GRAY,
            'תמצית תפוסה':     _TAB_CAPACITY,
            'תפוסה לפי אזור':  _TAB_TEAL,
            'תפוסה לפי בניין': _TAB_LIGHT_BLUE,
            'תפוסה לפי דירה':  _TAB_GREEN,
            'תפוסה לפי חדרים': _TAB_FULL,
        }
        for ws in book.worksheets:
            _format_worksheet(ws)
            _set_tab_color(ws, tab_colors.get(ws.title, _TAB_CAPACITY))

        for sheet_name in ['תפוסה לפי אזור', 'תפוסה לפי בניין', 'תפוסה לפי דירה', 'תפוסה לפי חדרים']:
            _color_occupancy_rows(book[sheet_name])

        _reapply_headers(book)

    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# Report 4: Manual Review Report
# ---------------------------------------------------------------------------

def generate_manual_review_report(region_id=None, region_name=None):
    d          = _load_data(region_id=region_id)
    region_lbl = region_name or ('כל האזורים' if not region_id else region_id)

    expl_lines = [
        'דוח בדיקה ידנית – Dormify',
        '',
        f'אזור נבחר: {region_lbl}',
        '',
        'הדוח מכיל רק רשומות הדורשות בדיקה ידנית.',
        'אינו רשימת סטודנטים כללית – מתמקד בחריגים ובנתונים חסרים/בעייתיים.',
        '',
        'גיליון "חריגים לבדיקה"         – כל הרשומות הבעייתיות במקום אחד.',
        'גיליון "תעודות זהות כפולות"    – ת.ז. שמופיעה ביותר מרשומה אחת.',
        'גיליון "נתונים חסרים"           – שדות חסרים: שם / מגדר / ת.ז. / דיור.',
        'גיליון "לא שובצו עם בעיה"       – לא משובצים שיש להם גם חריגת נתונים.',
    ]

    # Computed once, in a single pass, and reused by the 3 sheets below
    # that need it (summary counts, full exceptions list, unassigned-with-
    # issues) - see _compute_student_issues(). _build_duplicate_id_rows
    # and _build_missing_data_rows compute genuinely different things
    # (duplicate-ID grouping; a missing-data-only check with no
    # duplicate-ID component) and are left as their own passes.
    student_issues = _compute_student_issues(d['all_students'])

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        _expl_df(expl_lines).to_excel(writer, sheet_name='הסבר', index=False)

        _to_df(_build_exceptions_summary_rows(student_issues, d['assignment_map'], region_lbl)
               ).to_excel(writer, sheet_name='תמצית חריגים', index=False)

        _to_df(_build_all_exceptions_rows(student_issues, d['assignment_map'])
               ).to_excel(writer, sheet_name='חריגים לבדיקה', index=False)

        _to_df(_build_duplicate_id_rows(d['all_students'])
               ).to_excel(writer, sheet_name='תעודות זהות כפולות', index=False)

        _to_df(_build_missing_data_rows(d['all_students'])
               ).to_excel(writer, sheet_name='נתונים חסרים', index=False)

        _to_df(_build_unassigned_with_issues(student_issues, d['assignment_map'])
               ).to_excel(writer, sheet_name='לא שובצו עם בעיה', index=False)

        book = writer.book
        tab_colors = {
            'הסבר':                 _TAB_GRAY,
            'תמצית חריגים':        _TAB_REVIEW,
            'חריגים לבדיקה':       _TAB_RED,
            'תעודות זהות כפולות':  _TAB_DARK_ORANGE,
            'נתונים חסרים':         _TAB_ORANGE,
            'לא שובצו עם בעיה':    _TAB_AMBER,
        }
        for ws in book.worksheets:
            _format_worksheet(ws)
            _set_tab_color(ws, tab_colors.get(ws.title, _TAB_REVIEW))

        _apply_uniform_row_fill(book['חריגים לבדיקה'],      _FILL_EXCEPTION)
        _apply_uniform_row_fill(book['תעודות זהות כפולות'], _FILL_DUP_ID)
        _apply_uniform_row_fill(book['נתונים חסרים'],        _FILL_MISSING)
        _apply_uniform_row_fill(book['לא שובצו עם בעיה'],   _FILL_UNASSIGNED)

        _reapply_headers(book)

    buffer.seek(0)
    return buffer
